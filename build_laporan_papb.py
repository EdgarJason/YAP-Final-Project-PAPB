import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os
import shutil

# ==========================================
# ATURAN TIPOGRAFI BAKU (JARNIK STANDARD)
# - 100% HITAM
# - NO ITALIC
# - BOLD HANYA PADA JUDUL DAN SUB JUDUL
# - HINDARI DALAM KURUNG YANG TIDAK PERLU
# ==========================================
FONT_NAME = "Times New Roman"
COLOR_BLACK = RGBColor(0, 0, 0)
LOGO_UB = r"C:\Users\kuahs\Documents\MLOPS\scratch\ghazy_template_media\image3.png"

def set_cell_background(cell, fill_hex="FFFFFF"):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def set_run_font(run, size=12, bold=False):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = False
    run.font.color.rgb = COLOR_BLACK
    return run

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4, space_before=0, line_spacing=1.15):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)

def add_body(doc, text):
    """Paragraf narasi biasa, tidak bold, rata kanan-kiri rapi."""
    p = doc.add_paragraph()
    format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4, space_before=0)
    r = p.add_run(text.strip())
    set_run_font(r, size=12, bold=False)
    return p

def add_h1(doc, text):
    """Judul bab / bagian utama, BOLD."""
    p = doc.add_paragraph()
    format_p(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6, space_before=14)
    r = p.add_run(text.strip())
    set_run_font(r, size=13, bold=True)
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))
    return p

def add_h2(doc, text):
    """Subjudul bagian, BOLD."""
    p = doc.add_paragraph()
    format_p(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4, space_before=8)
    r = p.add_run(text.strip())
    set_run_font(r, size=12, bold=True)
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))
    return p

def add_bullet(doc, prefix, text, level=0):
    """Item poin bernomor atau bullet, TIDAK BOLD."""
    p = doc.add_paragraph()
    format_p(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2, space_before=1)
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    p.paragraph_format.first_line_indent = Inches(-0.25)
    r_prefix = p.add_run(prefix + " ")
    set_run_font(r_prefix, size=12, bold=False)
    r_text = p.add_run(text.strip())
    set_run_font(r_text, size=12, bold=False)
    return p

def build_laporan():
    print("[*] Membangun dokumen laporan Tugas 1 PAPB standar baku Jarnik...")
    doc = docx.Document()

    # 1.0 inch standard margins
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # ==========================================
    # HALAMAN JUDUL & KEPALA DOKUMEN
    # ==========================================
    p_t1 = doc.add_paragraph()
    format_p(p_t1, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4, space_before=6)
    set_run_font(p_t1.add_run("TUGAS 1 - USER INTERFACE"), size=15, bold=True)

    p_t2 = doc.add_paragraph()
    format_p(p_t2, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4, space_before=2)
    set_run_font(p_t2.add_run("YAP MOBILE AI STUDIO: PERANCANGAN DAN SLICING ANTARMUKA PENGGUNA BERBASIS JETPACK COMPOSE"), size=12, bold=True)

    p_t3 = doc.add_paragraph()
    format_p(p_t3, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14, space_before=2)
    set_run_font(p_t3.add_run("Mata Kuliah: Pemrograman Aplikasi Perangkat Bergerak / 2026"), size=11, bold=True)

    # Logo UB bila tersedia
    if os.path.exists(LOGO_UB):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(4)
        p_logo.paragraph_format.space_after = Pt(14)
        r_logo = p_logo.add_run()
        r_logo.add_picture(LOGO_UB, width=Inches(1.8))

    p_c5 = doc.add_paragraph()
    format_p(p_c5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6, space_before=2)
    set_run_font(p_c5.add_run("Disusun Oleh:\nKELOMPOK 2"), size=11, bold=True)

    # Tabel Anggota Kelompok (Header tidak bold, tanpa kurung)
    tbl_author = doc.add_table(rows=1, cols=3)
    tbl_author.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_author.autofit = False
    tbl_author.columns[0].width = Inches(2.6)
    tbl_author.columns[1].width = Inches(1.8)
    tbl_author.columns[2].width = Inches(2.1)
    set_table_borders(tbl_author, color="000000", sz="4", val="single")

    hdr_cols = ["Nama Anggota", "NIM", "Peran dan Tanggung Jawab"]
    for c_i, h_txt in enumerate(hdr_cols):
        c = tbl_author.cell(0, c_i)
        set_cell_background(c, "E5E7EB")
        set_cell_margins(c, top=60, bottom=60, left=80, right=80)
        p = c.paragraphs[0]
        format_p(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, space_before=2)
        r = p.add_run(h_txt)
        set_run_font(r, size=10, bold=False)

    members = [
        ("Edgar Jason Husin", "245150201111047", "UI Designer dan Frontend Slicing"),
        ("Joshua Nathanael Purba", "245150201111001", "Frontend Slicing dan Interaktivitas"),
        ("Tobias Andra Valentino", "245150200111076", "Koordinator Tim dan Slicing Compose"),
        ("Willy Gregory", "245150201111089", "Validasi Desain dan Pengujian UI")
    ]

    for m_nama, m_nim, m_peran in members:
        row = tbl_author.add_row()
        for c_i, val in enumerate([m_nama, m_nim, m_peran]):
            c = row.cells[c_i]
            set_cell_margins(c, top=50, bottom=50, left=80, right=80)
            p = c.paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.CENTER if c_i == 1 else WD_ALIGN_PARAGRAPH.LEFT
            format_p(p, align=align, space_after=1, space_before=1)
            r = p.add_run(val)
            set_run_font(r, size=10, bold=False)

    p_sp = doc.add_paragraph()
    format_p(p_sp, space_after=10, space_before=0)

    # ==========================================
    # BAGIAN 1: INFORMASI DAN TAUTAN PROYEK
    # ==========================================
    add_h1(doc, "1. Informasi dan Tautan Proyek")

    add_body(doc, "Seluruh sumber daya pengembangan antarmuka pengguna aplikasi YAP Mobile AI Studio telah diunggah dan dapat diakses melalui tautan repositori kode dan desain berikut:")

    tbl_links = doc.add_table(rows=1, cols=3)
    tbl_links.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_links.autofit = False
    tbl_links.columns[0].width = Inches(1.8)
    tbl_links.columns[1].width = Inches(3.2)
    tbl_links.columns[2].width = Inches(1.5)
    set_table_borders(tbl_links, color="000000", sz="4", val="single")

    hdr_links = ["Komponen Proyek", "Tautan Resmi", "Keterangan"]
    for c_i, h_txt in enumerate(hdr_links):
        c = tbl_links.cell(0, c_i)
        set_cell_background(c, "E5E7EB")
        set_cell_margins(c, top=60, bottom=60, left=80, right=80)
        p = c.paragraphs[0]
        format_p(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, space_before=2)
        r = p.add_run(h_txt)
        set_run_font(r, size=10, bold=False)

    link_data = [
        ("Repositori GitHub", "https://github.com/EdgarJason/YAP-Final-Project-PAPB", "Kode Sumber Jetpack Compose"),
        ("Rancangan Figma", "https://www.figma.com/design/odHg0ZJ4Vu64v4RRieNj7U/YAP---Final-Project-PAPB?node-id=20-379&t=NHgl5ypcp4RNecBd-1", "Master Desain UI dan Alur Navigasi")
    ]

    for k, u, desc in link_data:
        row = tbl_links.add_row()
        for c_i, val in enumerate([k, u, desc]):
            c = row.cells[c_i]
            set_cell_margins(c, top=50, bottom=50, left=80, right=80)
            p = c.paragraphs[0]
            format_p(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=1, space_before=1)
            r = p.add_run(val)
            set_run_font(r, size=9.5, bold=False)

    # ==========================================
    # BAGIAN 2: OVERVIEW PROYEK
    # ==========================================
    add_h1(doc, "2. Overview Proyek")

    add_body(doc, "YAP Mobile AI Studio merupakan aplikasi perangkat bergerak berbasis sistem operasi Android yang dirancang untuk membantu pengguna menghasilkan prototipe aplikasi berbasis web dengan memanfaatkan kecerdasan buatan atau Artificial Intelligence. Aplikasi ini memungkinkan pengguna menyampaikan kebutuhan dan spesifikasi pembuatan aplikasi secara interaktif melalui percakapan alami, kemudian menelaah kode program yang dihasilkan untuk menyempurnakan aplikasi web sesuai kebutuhan pengguna. Target pengguna aplikasi ini mencakup mahasiswa, pelajar, pengembang perangkat lunak tingkat awal, serta masyarakat umum yang memerlukan sarana perancangan web cepat tanpa keharusan menggunakan komputer meja secara terus-menerus.")

    add_body(doc, "Arsitektur antarmuka YAP Mobile AI Studio dibangun dengan mengedepankan beberapa fitur inti yang saling terintegrasi:")

    fitur_list = [
        ("1.", "Autentikasi Pengguna: Mengelola pendaftaran akun baru, proses masuk, verifikasi kode One-Time Password, serta pemulihan kata sandi."),
        ("2.", "Profil Pengguna: Memfasilitasi pengelolaan informasi akun seperti pembaruan nama pengguna, alamat email, kata sandi, serta penonaktifan akun."),
        ("3.", "Chat AI Studio: Ruang kerja utama komunikasi pengguna dengan mesin kecerdasan buatan guna menyampaikan instruksi pembuatan dan pembaruan kode aplikasi."),
        ("4.", "Code Inspector: Penampil kode program hasil generasi AI yang memudahkan pengguna memeriksa struktur kode HTML, CSS, maupun JavaScript."),
        ("5.", "Pengaturan API: Konfigurasi antarmuka pemrograman aplikasi pihak ketiga untuk menghubungkan aplikasi dengan penyedia model kecerdasan buatan."),
        ("6.", "Pratinjau Web: Fitur penampil tampilan web secara langsung sehingga hasil rancangan dapat dievaluasi sebelum diunduh atau dipublikasikan."),
        ("7.", "Riwayat Percakapan: Penyimpanan arsip instruksi dan percakapan lampau agar proses iterasi pengembangan dapat dilanjutkan kapan saja.")
    ]

    for pfx, txt in fitur_list:
        add_bullet(doc, pfx, txt)

    add_body(doc, "Integrasi seluruh fungsi tersebut ke dalam satu aplikasi bergerak memberikan alur kerja perancangan web yang lebih fleksibel, portabel, dan mudah dipahami oleh pengguna pemula.")

    # ==========================================
    # BAGIAN 3: TAMPILAN ANTARMUKA PENGGUNA (HASIL SLICING UI)
    # ==========================================
    add_h1(doc, "3. Tampilan Antarmuka Pengguna Hasil Slicing")

    add_body(doc, "Seluruh rancangan antarmuka pengguna pada Figma telah diimplementasikan ke dalam komponen Jetpack Compose pada Android Studio. Tabel berikut menyajikan tangkapan layar hasil slicing beserta penjelasan fungsi tiap antarmuka:")

    screens_data = [
        ("1", "extracted_images/screen_1_Welcome_screen.png", "Welcome Screen",
         "Layar selamat datang yang pertama kali ditampilkan kepada pengguna. Menampilkan identitas visual YAP, maskot speech-bubble yang ramah, serta tombol navigasi utama menuju Log In dan Sign Up."),
        ("2", "extracted_images/screen_2_Registe_screen.png", "Register Screen",
         "Layar pendaftaran akun baru. Menyediakan formulir input nama lengkap, email, kata sandi, dan konfirmasi kata sandi dengan validasi interaktif sebelum akun dibuat."),
        ("3", "extracted_images/screen_3_Profile_screen.png", "Profile Screen",
         "Layar profil akun pengguna. Menampilkan informasi akun, avatar pengguna, opsi navigasi ubah nama, ganti email, perbarui kata sandi, tombol logout, serta opsi hapus akun."),
        ("4", "extracted_images/screen_4_OTP_screen.png", "OTP Verification Screen",
         "Layar verifikasi keamanan berupa kode One-Time Password empat digit yang dikirimkan ke email pengguna untuk keperluan autentikasi ganda atau pemulihan sandi."),
        ("5", "extracted_images/screen_5_Login_screen.png", "Login Screen",
         "Layar autentikasi masuk pengguna terdaftar. Memuat kolom isian email, kata sandi, opsi Remember Me, tautan bantuan lupa kata sandi, serta tombol masuk."),
        ("6", "extracted_images/screen_6_Forgot_password_screen.png", "Forgot Password Screen",
         "Layar permintaan pemulihan kata sandi. Pengguna memasukkan alamat email terdaftar untuk menerima tautan atau kode reset kata sandi."),
        ("7", "extracted_images/screen_7_Delete_account_screen.png", "Delete Account Screen",
         "Layar konfirmasi penghapusan akun permanen. Dilengkapi peringatan konsekuensi kehilangan data dan konfirmasi sandi untuk mencegah tindakan yang tidak disengaja."),
        ("8", "extracted_images/screen_8_Chat_studio_screen.png", "Chat Studio Screen",
         "Layar kerja utama studio AI. Menyediakan bilah drawer ruang kerja, kolom masukan instruksi atau ide web, tombol pengiriman prompt, dan tabs navigasi antarmuka studio."),
        ("9", "extracted_images/screen_9_Chat_history_screen.png", "Chat History Screen",
         "Layar arsip percakapan terdahulu. Memuat daftar riwayat interaksi pengguna dengan AI, fitur pencarian riwayat percakapan, serta pintasan memulai sesi baru."),
        ("10", "extracted_images/screen_10_Change_password_screen.png", "Change Password Screen",
         "Layar pembaruan kata sandi pengguna. Memerlukan input kata sandi saat ini, kata sandi baru, dan konfirmasi kata sandi baru untuk menjaga integritas akun."),
        ("11", "extracted_images/screen_11_Change_name_screen.png", "Change Name Screen",
         "Layar penyuntingan nama profil pengguna. Menyediakan kolom pembaruan nama tampilan yang langsung tersimpan ke sistem profil."),
        ("12", "extracted_images/screen_12_Change_email_screen.png", "Change Email Screen",
         "Layar pembaruan alamat email akun. Memfasilitasi pergantian email dengan protokol verifikasi keamanan ke email baru."),
        ("13", "extracted_images/screen_13_API_settings_screen.png", "API Settings Screen",
         "Layar konfigurasi penyedia layanan kecerdasan buatan. Memungkinkan pengguna mengatur kunci API dan endpoint model guna menghubungkan studio dengan AI pihak ketiga."),
        ("14", "extracted_images/screen_14_All_set_screen.png", "All Set Screen",
         "Layar konfirmasi penyelesaian registrasi dan verifikasi. Menandakan bahwa akun telah terkonfigurasi secara penuh dan pengguna siap menjelajahi studio pembuatan web.")
    ]

    tbl_screens = doc.add_table(rows=1, cols=3)
    tbl_screens.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_screens.autofit = False
    tbl_screens.columns[0].width = Inches(0.6)
    tbl_screens.columns[1].width = Inches(2.4)
    tbl_screens.columns[2].width = Inches(3.5)
    set_table_borders(tbl_screens, color="000000", sz="4", val="single")

    hdr_sc = ["No.", "Tangkapan Layar", "Keterangan dan Deskripsi Antarmuka"]
    for c_i, h_txt in enumerate(hdr_sc):
        c = tbl_screens.cell(0, c_i)
        set_cell_background(c, "E5E7EB")
        set_cell_margins(c, top=60, bottom=60, left=80, right=80)
        p = c.paragraphs[0]
        format_p(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, space_before=2)
        r = p.add_run(h_txt)
        set_run_font(r, size=10, bold=False)

    for no_str, img_path, title_sc, desc_sc in screens_data:
        row = tbl_screens.add_row()

        # Kolom No
        c_no = row.cells[0]
        set_cell_margins(c_no, top=60, bottom=60, left=60, right=60)
        p_no = c_no.paragraphs[0]
        format_p(p_no, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0, space_before=0)
        r_no = p_no.add_run(no_str)
        set_run_font(r_no, size=10, bold=False)

        # Kolom Gambar
        c_img = row.cells[1]
        set_cell_margins(c_img, top=60, bottom=60, left=60, right=60)
        p_img = c_img.paragraphs[0]
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(4)
        if os.path.exists(img_path):
            r_img = p_img.add_run()
            r_img.add_picture(img_path, width=Inches(2.0))

        # Kolom Keterangan
        c_desc = row.cells[2]
        set_cell_margins(c_desc, top=60, bottom=60, left=80, right=80)
        p_title = c_desc.paragraphs[0]
        format_p(p_title, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=3, space_before=2)
        r_t = p_title.add_run(title_sc)
        set_run_font(r_t, size=10.5, bold=True)

        p_desc = c_desc.add_paragraph()
        format_p(p_desc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=2, space_before=0)
        r_d = p_desc.add_run(desc_sc)
        set_run_font(r_d, size=10, bold=False)

    # ==========================================
    # BAGIAN 4: CATATAN KOMPONEN DAN NAVIGASI
    # ==========================================
    add_h1(doc, "4. Catatan Komponen dan Navigasi Antarmuka")

    add_body(doc, "Komponen navigasi seperti Modal Navigation Drawer, pop-up dialog konfirmasi, dan bilah input dirancang secara reusable pada paket komponen Jetpack Compose. Implementasi detail kode interaksi dan animasi responsif dapat ditinjau langsung pada berkas repositori GitHub, serta visualisasi alur perpindahan layar dapat diperiksa secara komprehensif pada dokumen rancangan Figma.")

    # Simpan Berkas
    out_docx = "KELOMPOK_2_TUGAS_1_USER_INTERFACE.docx"
    doc.save(out_docx)
    print(f"[SUCCESS] Berkas berhasil disimpan ke: {out_docx}")

    # Salin ke Google Drive Desktop jika G:\Drive Saya ada
    gdrive_path = r"G:\Drive Saya\KELOMPOK 2_TUGAS 1-USER INTERFACE.docx"
    try:
        if os.path.exists(r"G:\Drive Saya"):
            shutil.copyfile(out_docx, gdrive_path)
            print(f"[SUCCESS] Berhasil disinkronkan ke Google Drive: {gdrive_path}")
    except Exception as e:
        print(f"[WARNING] Gagal menyalin ke Google Drive: {e}")

if __name__ == "__main__":
    build_laporan()
