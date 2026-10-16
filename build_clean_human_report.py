import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
import os
import shutil
import win32com.client

# ==============================================================================
# STANDAR MUTLAK LAPORAN ANTI-AI SLOP (FILKOM UB / STANDAR JARNIK & PAPB)
# ==============================================================================
# 1. Tipografi 100% Konsisten:
#    - Font: Times New Roman di seluruh dokumen tanpa terkecuali.
#    - Warna teks: 100% Hitam Murni (RGB 0,0,0 / #000000).
#    - Tanpa Italic (No Italic).
#    - BOLD HANYA untuk Judul Dokumen, Label Bagian Utama, dan Header Tabel.
#    - Isi teks narasi, link, nama anggota, dan isi sel tabel DILARANG DI-BOLD.
#
# 2. Eliminasi Total "AI Slop" & Fabrikasi Teks:
#    - TIDAK ADA kalimat penjelas fiktif buatan AI pada tabel (kolom keterangan
#      hanya memuat nama layar asli dari mahasiswa, tanpa narasi karangan chatbot).
#    - TIDAK ADA penambahan tabel ekstra / sekat berlebihan yang tidak diminta.
#    - TIDAK ADA kata-kata klise korporat AI ("delve", "pivotal", "seamless", dll).
#    - Mempertahankan suara tulisan asli mahasiswa (Edgar Jason Husin, Joshua,
#      Tobias, Willy), hanya memperbaiki typo minor ("memanfaat" -> "memanfaatkan",
#      "yang ini membuat" -> "yang ingin membuat", "Registe" -> "Register").
#
# 3. Struktur Tabel & Layouting Bersih:
#    - Header tabel abu-abu terang (#E5E7EB) dengan border tunggal hitam (0.5 pt).
#    - Padding sel internal eksplisit (top/bottom 60 dxa, left/right 80 dxa).
#    - Properti XML <w:cantSplit/> pada setiap baris agar tidak terbelah canggung.
#    - Properti XML <w:tblHeader/> pada baris header agar konsisten antar halaman.
#    - Perataan vertikal tengah (w:vAlign val="center") pada sel tabel.
#    - Menghindari "Justification Stretch Bug" (Justify hanya untuk paragraf panjang).
# ==============================================================================

FONT_NAME = "Times New Roman"
COLOR_BLACK = RGBColor(0, 0, 0)

def set_cell_background(cell, fill_hex="FFFFFF"):
    tcPr = cell._element.get_or_add_tcPr()
    for old_shd in tcPr.findall(qn('w:shd')):
        tcPr.remove(old_shd)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._element.get_or_add_tcPr()
    for old_mar in tcPr.findall(qn('w:tcMar')):
        tcPr.remove(old_mar)
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_valign(cell, val="center"):
    tcPr = cell._element.get_or_add_tcPr()
    for old_v in tcPr.findall(qn('w:vAlign')):
        tcPr.remove(old_v)
    vAlign = parse_xml(f'<w:vAlign {nsdecls("w")} w:val="{val}"/>')
    tcPr.append(vAlign)

def set_table_borders(table, color="000000", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        for old_b in tblPr[0].findall(qn('w:tblBorders')):
            tblPr[0].remove(old_b)
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def set_run(run, size=12, bold=False):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = False
    run.font.color.rgb = COLOR_BLACK

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4, space_before=0, line_spacing=1.15):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)

def build_report():
    print("[*] Membangun dokumen Tugas 1 PAPB murni tanpa AI slop...")
    doc = docx.Document()

    # Margin standar akademik 1.0 inci di semua sisi
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # 1. JUDUL DOKUMEN (Centered, Bold, 14 pt)
    p_title = doc.add_paragraph()
    format_p(p_title, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14, space_before=0)
    set_run(p_title.add_run("TUGAS 1 - USER INTERFACE"), size=14, bold=True)

    # 2. IDENTITAS KELOMPOK
    p_kel = doc.add_paragraph()
    format_p(p_kel, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2, space_before=0)
    set_run(p_kel.add_run("Kelompok 2:"), size=12, bold=True)

    members = [
        "1. Edgar Jason Husin - 245150201111047",
        "2. Joshua Nathanael Purba - 245150201111001",
        "3. Tobias Andra Valentino - 245150200111076",
        "4. Willy Gregory - 245150201111089"
    ]
    for m in members:
        p_m = doc.add_paragraph()
        format_p(p_m, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2, space_before=0)
        p_m.paragraph_format.left_indent = Inches(0.2)
        set_run(p_m.add_run(m), size=12, bold=False)

    p_spacer1 = doc.add_paragraph()
    p_spacer1.paragraph_format.space_before = Pt(0)
    p_spacer1.paragraph_format.space_after = Pt(4)

    # 3. TAUTAN PROYEK (Link GitHub & Figma)
    p_gh_label = doc.add_paragraph()
    format_p(p_gh_label, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=1, space_before=0)
    set_run(p_gh_label.add_run("Link GitHub Proyek:"), size=12, bold=True)

    p_gh = doc.add_paragraph()
    format_p(p_gh, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6, space_before=0)
    set_run(p_gh.add_run("https://github.com/EdgarJason/YAP-Final-Project-PAPB"), size=11, bold=False)

    p_fg_label = doc.add_paragraph()
    format_p(p_fg_label, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=1, space_before=0)
    set_run(p_fg_label.add_run("Link Figma:"), size=12, bold=True)

    p_fg = doc.add_paragraph()
    format_p(p_fg, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=12, space_before=0)
    set_run(p_fg.add_run("https://www.figma.com/design/odHg0ZJ4Vu64v4RRieNj7U/YAP---Final-Project-PAPB?node-id=20-379&t=NHgl5ypcp4RNecBd-1"), size=11, bold=False)

    # 4. OVERVIEW PROYEK (Narasi Otentik Mahasiswa)
    p_ov_label = doc.add_paragraph()
    format_p(p_ov_label, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4, space_before=4)
    set_run(p_ov_label.add_run("OVERVIEW PROYEK:"), size=12, bold=True)
    pPr_ov = p_ov_label._p.get_or_add_pPr()
    pPr_ov.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))

    # Paragraf 1
    p_ov1 = doc.add_paragraph()
    format_p(p_ov1, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, space_before=0)
    set_run(p_ov1.add_run(
        "YAP Mobile AI Studio merupakan aplikasi perangkat bergerak berbasis Android yang dirancang untuk membantu pengguna "
        "menghasilkan aplikasi berbasis web dengan memanfaatkan Artificial Intelligence (AI). Aplikasi ini memungkinkan pengguna "
        "menyampaikan kebutuhan pengembangan aplikasi melalui percakapan dengan AI, kemudian memanfaatkan kode program yang "
        "dihasilkan untuk membangun dan menyempurnakan aplikasi web sesuai dengan kebutuhan. Target pengguna aplikasi ini meliputi "
        "mahasiswa, pelajar, pengembang perangkat lunak pemula, serta pengguna lain yang ingin membuat aplikasi berbasis web dengan "
        "bantuan AI. Melalui antarmuka mobile, pengguna diharapkan dapat mengakses berbagai fungsi pengembangan aplikasi secara lebih "
        "praktis tanpa harus selalu menggunakan komputer."
    ), size=12, bold=False)

    # Paragraf 2
    p_ov2 = doc.add_paragraph()
    format_p(p_ov2, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, space_before=0)
    set_run(p_ov2.add_run(
        "Aplikasi YAP Mobile AI Studio memiliki beberapa fitur utama yang saling mendukung proses pengembangan aplikasi. "
        "Fitur authentication atau auth digunakan untuk mengelola proses autentikasi dan akses pengguna ke dalam aplikasi, "
        "sedangkan fitur profile memungkinkan pengguna melihat dan mengelola informasi profil akun. Fitur chat menjadi sarana "
        "komunikasi antara pengguna dan AI untuk memberikan instruksi pembuatan aplikasi, meminta perubahan kode, serta "
        "mengembangkan fitur sesuai kebutuhan. Hasil dari interaksi tersebut dapat ditinjau melalui fitur code yang menampilkan "
        "kode program hasil generasi AI sehingga pengguna dapat memeriksa dan mengelola kode yang digunakan aplikasi berbasis web."
    ), size=12, bold=False)

    # Paragraf 3
    p_ov3 = doc.add_paragraph()
    format_p(p_ov3, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=14, space_before=0)
    set_run(p_ov3.add_run(
        "Selain itu, YAP Mobile AI Studio menyediakan fitur API Settings untuk mengatur konfigurasi API yang diperlukan dalam "
        "menghubungkan aplikasi dengan layanan AI. Fitur preview memungkinkan pengguna melihat pratinjau tampilan aplikasi web yang "
        "dihasilkan sehingga hasil pengembangan dapat diperiksa dan dievaluasi sebelum dilakukan penyempurnaan lebih lanjut. "
        "Sementara itu, fitur chat history digunakan untuk menyimpan dan menampilkan riwayat percakapan antara pengguna dan AI, "
        "sehingga pengguna dapat mengakses kembali instruksi sebelumnya serta melanjutkan proses pengembangan berdasarkan "
        "percakapan yang telah dilakukan. Secara keseluruhan, YAP Mobile AI Studio dirancang untuk mengintegrasikan interaksi "
        "dengan AI, pengelolaan kode program, pengaturan API, pratinjau aplikasi, dan penyimpanan riwayat percakapan dalam satu "
        "platform mobile. Integrasi tersebut diharapkan dapat memberikan pengalaman pengembangan aplikasi berbasis web yang lebih "
        "praktis, fleksibel, dan mudah diakses oleh pengguna."
    ), size=12, bold=False)

    # 5. TAMPILAN UI (SUDAH DISLICING)
    p_ui_label = doc.add_paragraph()
    format_p(p_ui_label, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6, space_before=4)
    set_run(p_ui_label.add_run("TAMPILAN UI (SUDAH DISLICING)"), size=12, bold=True)
    pPr_ui = p_ui_label._p.get_or_add_pPr()
    pPr_ui.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))

    # 14 Tangkapan Layar Murni Sesuai Dokumen Asli Mahasiswa (Tanpa Deskripsi Tambahan AI Slop)
    screens = [
        ("1", "extracted_images/screen_1_Welcome_screen.png", "Welcome screen"),
        ("2", "extracted_images/screen_2_Registe_screen.png", "Register screen"),
        ("3", "extracted_images/screen_3_Profile_screen.png", "Profile screen"),
        ("4", "extracted_images/screen_4_OTP_screen.png", "OTP screen"),
        ("5", "extracted_images/screen_5_Login_screen.png", "Login screen"),
        ("6", "extracted_images/screen_6_Forgot_password_screen.png", "Forgot password screen"),
        ("7", "extracted_images/screen_7_Delete_account_screen.png", "Delete account screen"),
        ("8", "extracted_images/screen_8_Chat_studio_screen.png", "Chat studio screen"),
        ("9", "extracted_images/screen_9_Chat_history_screen.png", "Chat history screen"),
        ("10", "extracted_images/screen_10_Change_password_screen.png", "Change password screen"),
        ("11", "extracted_images/screen_11_Change_name_screen.png", "Change name screen"),
        ("12", "extracted_images/screen_12_Change_email_screen.png", "Change email screen"),
        ("13", "extracted_images/screen_13_API_settings_screen.png", "API settings screen"),
        ("14", "extracted_images/screen_14_All_set_screen.png", "All set screen")
    ]

    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(0.6)
    tbl.columns[1].width = Inches(2.7)
    tbl.columns[2].width = Inches(3.2)
    set_table_borders(tbl, color="000000", sz="4", val="single")

    # Header Tabel: Abu-abu terang (#E5E7EB), padding dxa, cantSplit & tblHeader
    hdr = tbl.rows[0]
    trPr = hdr._element.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    headers = ["No.", "Gambar", "Keterangan"]
    for c_idx, h_txt in enumerate(headers):
        cell = hdr.cells[c_idx]
        set_cell_background(cell, "E5E7EB")
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        set_cell_valign(cell, "center")
        p = cell.paragraphs[0]
        format_p(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1, space_before=1)
        set_run(p.add_run(h_txt), size=11, bold=True)

    # Baris Isi Data
    for no, img_p, title in screens:
        row = tbl.add_row()
        trPr_r = row._element.get_or_add_trPr()
        trPr_r.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        # 1. Sel Nomor (Center)
        c_no = row.cells[0]
        set_cell_background(c_no, "FFFFFF")
        set_cell_margins(c_no, top=60, bottom=60, left=60, right=60)
        set_cell_valign(c_no, "center")
        p_n = c_no.paragraphs[0]
        format_p(p_n, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0, space_before=0)
        set_run(p_n.add_run(no), size=11, bold=False)

        # 2. Sel Gambar (Center)
        c_img = row.cells[1]
        set_cell_background(c_img, "FFFFFF")
        set_cell_margins(c_img, top=60, bottom=60, left=60, right=60)
        set_cell_valign(c_img, "center")
        p_i = c_img.paragraphs[0]
        p_i.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_i.paragraph_format.space_before = Pt(4)
        p_i.paragraph_format.space_after = Pt(4)
        if os.path.exists(img_p):
            r_img = p_i.add_run()
            r_img.add_picture(img_p, width=Inches(2.2))

        # 3. Sel Keterangan (Left, Vertically Centered, Murni Nama Layar Tanpa Bolding Berlebih)
        c_desc = row.cells[2]
        set_cell_background(c_desc, "FFFFFF")
        set_cell_margins(c_desc, top=60, bottom=60, left=100, right=80)
        set_cell_valign(c_desc, "center")
        p_t = c_desc.paragraphs[0]
        format_p(p_t, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0, space_before=0)
        set_run(p_t.add_run(title), size=11.5, bold=False)

    # 6. CATATAN PENUTUP
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    p_foot = doc.add_paragraph()
    format_p(p_foot, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4, space_before=6)
    set_run(p_foot.add_run("Navigasi dan Popup dapat dilihat pada GitHub, karena merupakan komponen. Dapat dilihat pula selengkapnya di Figma."), size=11, bold=False)

    # Simpan DOCX
    out_docx = "KELOMPOK_2_TUGAS_1_USER_INTERFACE.docx"
    doc.save(out_docx)
    print(f"[OK] Berkas DOCX murni berhasil dibuat: {out_docx}")

    # Salin ke Google Drive Desktop
    gdrive_docx = r"G:\Drive Saya\KELOMPOK 2_TUGAS 1-USER INTERFACE.docx"
    if os.path.exists(r"G:\Drive Saya"):
        shutil.copy2(out_docx, gdrive_docx)
        print(f"[OK] Berhasil PUSH DOCX ke Google Drive: {gdrive_docx}")

    # Ekspor ke PDF via Microsoft Word COM
    print("[*] Mengonversi ke PDF via Word COM headless...")
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    try:
        abs_in = os.path.abspath(out_docx)
        abs_pdf = os.path.abspath("KELOMPOK_2_TUGAS_1_USER_INTERFACE.pdf")
        doc_w = word.Documents.Open(abs_in)
        doc_w.SaveAs(abs_pdf, FileFormat=17) # 17 = wdFormatPDF
        doc_w.Close()
        print(f"[OK] Berkas PDF berhasil dibuat: {abs_pdf}")

        gdrive_pdf = r"G:\Drive Saya\KELOMPOK 2_TUGAS 1-USER INTERFACE.pdf"
        if os.path.exists(r"G:\Drive Saya"):
            shutil.copy2(abs_pdf, gdrive_pdf)
            print(f"[OK] Berhasil PUSH PDF ke Google Drive: {gdrive_pdf}")
    except Exception as e:
        print(f"[WARNING] Gagal ekspor PDF: {e}")
    finally:
        word.Quit()

if __name__ == "__main__":
    build_report()
