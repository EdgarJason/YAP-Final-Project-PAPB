import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os
import shutil
import win32com.client

# ==========================================
# ATURAN TIPOGRAFI BAKU STANDAR JARNIK
# - 100% HITAM (RGB 0,0,0)
# - NO ITALIC
# - BOLD HANYA PADA JUDUL DAN SUB JUDUL
# - HINDARI DALAM KURUNG YANG TIDAK PERLU
# - TABLE ALIGNMENT: CENTER
# - CELL PADDING & BORDER CLEAN
# ==========================================
FONT_NAME = "Times New Roman"
COLOR_BLACK = RGBColor(0, 0, 0)
LOGO_UB = r"C:\Users\kuahs\Documents\MLOPS\scratch\ghazy_template_media\image3.png"

def set_cell_background(cell, fill_hex="FFFFFF"):
    tcPr = cell._element.get_or_add_tcPr()
    # Remove existing shd if any
    for old_shd in tcPr.findall(qn('w:shd')):
        tcPr.remove(old_shd)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    tcPr = cell._element.get_or_add_tcPr()
    for old_mar in tcPr.findall(qn('w:tcMar')):
        tcPr.remove(old_mar)
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        for old_b in tblPr[0].findall(qn('w:tblBorders')):
            tblPr[0].remove(old_b)
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def make_row_header(row):
    trPr = row._element.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def prevent_row_split(row):
    trPr = row._element.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def set_run_clean(run, size=12, bold=False):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = False
    run.font.color.rgb = COLOR_BLACK

def format_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4, space_before=0, line_spacing=1.15):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)

def clean_entire_document(doc_path, out_docx_path):
    print(f"[*] Memuat berkas: {doc_path}")
    doc = docx.Document(doc_path)

    # 1. Pastikan batas halaman standar 1.0 inch (Academic Standard)
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # 2. Format ulang seluruh paragraf
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            # Spacing minimal untuk baris kosong
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            continue

        # Deteksi Judul Utama
        if txt == "TUGAS 1 - USER INTERFACE":
            format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4, space_before=6)
            for r in p.runs:
                set_run_clean(r, size=15, bold=True)
        elif "YAP MOBILE AI STUDIO: PERANCANGAN" in txt:
            format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4, space_before=2)
            for r in p.runs:
                set_run_clean(r, size=12, bold=True)
        elif "Mata Kuliah: Pemrograman Aplikasi Perangkat Bergerak" in txt:
            format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14, space_before=2)
            for r in p.runs:
                set_run_clean(r, size=11, bold=True)
        elif txt.startswith("Disusun Oleh:"):
            format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6, space_before=2)
            for r in p.runs:
                set_run_clean(r, size=11, bold=True)
        # Deteksi Heading Bagian (1. , 2. , 3. , 4. )
        elif any(txt.startswith(f"{num}. ") for num in range(1, 10)):
            format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6, space_before=14)
            for r in p.runs:
                set_run_clean(r, size=13, bold=True)
            pPr = p._p.get_or_add_pPr()
            pPr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))
        # Deteksi Poin Bernomor Fitur di Overview (1. Autentikasi, 2. Profil, dst)
        elif any(txt.startswith(f"{num}. Autentikasi") or txt.startswith(f"{num}. Profil") or 
                 txt.startswith(f"{num}. Chat AI") or txt.startswith(f"{num}. Code") or 
                 txt.startswith(f"{num}. Pengaturan") or txt.startswith(f"{num}. Pratinjau") or 
                 txt.startswith(f"{num}. Riwayat") for num in range(1, 10)):
            format_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=2, space_before=1)
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            for r in p.runs:
                set_run_clean(r, size=11.5, bold=False)
        # Paragraf Isi Narasi
        else:
            format_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4, space_before=0)
            for r in p.runs:
                set_run_clean(r, size=12, bold=False)

    # 3. Format seluruh tabel standar akademik
    for t_idx, table in enumerate(doc.tables):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="000000", sz="4", val="single")

        # Table 0: Anggota Kelompok
        if t_idx == 0:
            col_widths = [Inches(2.5), Inches(1.8), Inches(2.2)]
            for row_idx, row in enumerate(table.rows):
                if row_idx == 0:
                    make_row_header(row)
                else:
                    prevent_row_split(row)
                for col_idx, cell in enumerate(row.cells):
                    cell.width = col_widths[col_idx]
                    if row_idx == 0:
                        set_cell_background(cell, "E5E7EB")
                        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                        for p in cell.paragraphs:
                            format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1, space_before=1)
                            for r in p.runs:
                                set_run_clean(r, size=10, bold=False)
                    else:
                        set_cell_background(cell, "FFFFFF")
                        set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
                        for p in cell.paragraphs:
                            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx == 1 else WD_ALIGN_PARAGRAPH.LEFT
                            format_paragraph(p, align=align, space_after=1, space_before=1)
                            for r in p.runs:
                                set_run_clean(r, size=10, bold=False)

        # Table 1: Tautan Proyek
        elif t_idx == 1:
            col_widths = [Inches(1.8), Inches(3.2), Inches(1.5)]
            for row_idx, row in enumerate(table.rows):
                if row_idx == 0:
                    make_row_header(row)
                else:
                    prevent_row_split(row)
                for col_idx, cell in enumerate(row.cells):
                    cell.width = col_widths[col_idx]
                    if row_idx == 0:
                        set_cell_background(cell, "E5E7EB")
                        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                        for p in cell.paragraphs:
                            format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1, space_before=1)
                            for r in p.runs:
                                set_run_clean(r, size=10, bold=False)
                    else:
                        set_cell_background(cell, "FFFFFF")
                        set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
                        for p in cell.paragraphs:
                            format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=1, space_before=1)
                            for r in p.runs:
                                set_run_clean(r, size=9.5, bold=False)

        # Table 2: Tangkapan Layar Antarmuka Slicing
        elif t_idx == 2:
            col_widths = [Inches(0.6), Inches(2.4), Inches(3.5)]
            for row_idx, row in enumerate(table.rows):
                if row_idx == 0:
                    make_row_header(row)
                else:
                    prevent_row_split(row)
                for col_idx, cell in enumerate(row.cells):
                    cell.width = col_widths[col_idx]
                    if row_idx == 0:
                        set_cell_background(cell, "E5E7EB")
                        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                        for p in cell.paragraphs:
                            format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1, space_before=1)
                            for r in p.runs:
                                set_run_clean(r, size=10, bold=False)
                    else:
                        set_cell_background(cell, "FFFFFF")
                        # Kolom No.
                        if col_idx == 0:
                            set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
                            for p in cell.paragraphs:
                                format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0, space_before=0)
                                for r in p.runs:
                                    set_run_clean(r, size=10, bold=False)
                        # Kolom Gambar
                        elif col_idx == 1:
                            set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
                            for p in cell.paragraphs:
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                p.paragraph_format.space_before = Pt(4)
                                p.paragraph_format.space_after = Pt(4)
                        # Kolom Keterangan
                        elif col_idx == 2:
                            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                            # Pisahkan judul layar dan deskripsi jika tergabung
                            full_text = cell.text.strip()
                            lines = [ln.strip() for ln in full_text.split('\n') if ln.strip()]
                            # Kosongkan paragraf lama lalu susun ulang rapi
                            cell.text = ""
                            if lines:
                                # Baris pertama: Judul Layar (Bold)
                                p_tit = cell.paragraphs[0]
                                format_paragraph(p_tit, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=3, space_before=2)
                                r_t = p_tit.add_run(lines[0])
                                set_run_clean(r_t, size=10.5, bold=True)
                                # Baris berikutnya: Deskripsi Layar (Normal Justified)
                                if len(lines) > 1:
                                    p_desc = cell.add_paragraph()
                                    format_paragraph(p_desc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=2, space_before=0)
                                    r_d = p_desc.add_run(" ".join(lines[1:]))
                                    set_run_clean(r_d, size=10, bold=False)

    doc.save(out_docx_path)
    print(f"[SUCCESS] Berkas rapi disimpan ke: {out_docx_path}")

def sync_to_gdrive_and_export_pdf(in_docx_path):
    # 1. Push ke Google Drive Desktop
    gdrive_docx = r"G:\Drive Saya\KELOMPOK 2_TUGAS 1-USER INTERFACE.docx"
    gdrive_pdf = r"G:\Drive Saya\KELOMPOK 2_TUGAS 1-USER INTERFACE.pdf"
    local_pdf = in_docx_path.replace(".docx", ".pdf")

    if os.path.exists(r"G:\Drive Saya"):
        shutil.copy2(in_docx_path, gdrive_docx)
        print(f"[SUCCESS] Berhasil PUSH ke Google Drive: {gdrive_docx}")

    # 2. Export ke PDF via Word COM
    print("[*] Mengonversi ke PDF secara headless...")
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    try:
        abs_in = os.path.abspath(in_docx_path)
        abs_out = os.path.abspath(local_pdf)
        doc = word.Documents.Open(abs_in)
        doc.SaveAs(abs_out, FileFormat=17) # 17 = wdFormatPDF
        doc.Close()
        print(f"[SUCCESS] PDF lokal berhasil dibuat: {abs_out}")

        if os.path.exists(r"G:\Drive Saya"):
            shutil.copy2(abs_out, gdrive_pdf)
            print(f"[SUCCESS] Berhasil PUSH PDF ke Google Drive: {gdrive_pdf}")
    except Exception as e:
        print(f"[WARNING] Gagal ekspor PDF: {e}")
    finally:
        word.Quit()

if __name__ == "__main__":
    pulled_doc = "laporan_latest_pulled.docx"
    refined_doc = "KELOMPOK_2_TUGAS_1_USER_INTERFACE.docx"
    clean_entire_document(pulled_doc, refined_doc)
    sync_to_gdrive_and_export_pdf(refined_doc)
